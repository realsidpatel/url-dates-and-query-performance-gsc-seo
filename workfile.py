# ------------------------------------------------------------
# 1. INSTALL / IMPORT LIBRARIES
# ------------------------------------------------------------

import re
import json
import requests
import pandas as pd

from bs4 import BeautifulSoup
from datetime import datetime

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# ------------------------------------------------------------
# 2. SETTINGS
# ------------------------------------------------------------

# Google Search Console date range
START_DATE = "2025-08-20"
END_DATE   = "2026-08-20"

# Maximum queries to retrieve per URL
MAX_QUERIES_PER_URL = 10

# Your GSC property
# Example:
# SITE_URL = "https://www.volopay.com/"
#
# IMPORTANT:
# Change this if your GSC property is different.
SITE_URL = "https://www.volopay.com/"

# Output filenames
XLSX_OUTPUT = "gsc_url_query_report.xlsx"
CSV_OUTPUT  = "gsc_url_query_report.csv"


# ------------------------------------------------------------
# 3. ENTER YOUR URLS
# ------------------------------------------------------------

urls = ["add a URL list here separated by comma"]

# ------------------------------------------------------------
# 4. DATE PARSING HELPERS
# ------------------------------------------------------------

def parse_date_value(value):
    """
    Convert a date/time string into a Python datetime.

    Handles examples such as:
      2024-04-05T13:24:00.670Z
      2024-12-06T07:34:13.512Z
      2024-04-05
      2024-04-05 13:24:00
    """

    if not value:
        return None

    value = str(value).strip()

    # Remove surrounding quotes
    value = value.strip('"').strip("'")

    # ISO UTC Z
    try:
        return datetime.fromisoformat(value.replace("Z", "+00:00"))
    except:
        pass

    # ISO without timezone
    formats = [
        "%Y-%m-%d",
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%dT%H:%M:%S",
        "%Y-%m-%dT%H:%M:%S.%f",
        "%Y/%m/%d",
        "%d-%m-%Y",
        "%d/%m/%Y",
    ]

    for fmt in formats:
        try:
            return datetime.strptime(value, fmt)
        except:
            continue

    return None


def format_date_only(dt):
    """
    Final report format:
    DD Month YYYY
    """

    if not dt:
        return ""

    return dt.strftime("%d %B %Y")


# ------------------------------------------------------------
# 5. EXTRACT DATE FROM PAGE SOURCE
# ------------------------------------------------------------

def extract_page_dates(url, timeout=30):

    headers = {
        "User-Agent": (
            "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
            "AppleWebKit/537.36 (KHTML, like Gecko) "
            "Chrome/131.0.0.0 Safari/537.36"
        )
    }

    result = {
        "created_datetime": None,
        "updated_datetime": None,
        "created_source": "",
        "updated_source": "",
        "status_code": None,
        "error": ""
    }

    try:
        response = requests.get(
            url,
            headers=headers,
            timeout=timeout
        )

        result["status_code"] = response.status_code

        if response.status_code != 200:
            result["error"] = f"HTTP {response.status_code}"
            return result

        source = response.text

        # ----------------------------------------------------
        # A. JSON-LD
        # ----------------------------------------------------

        soup = BeautifulSoup(source, "html.parser")

        json_ld_blocks = soup.find_all(
            "script",
            type="application/ld+json"
        )

        json_ld_objects = []

        for block in json_ld_blocks:

            try:
                data = json.loads(block.string or block.get_text())

                if isinstance(data, list):
                    json_ld_objects.extend(data)
                else:
                    json_ld_objects.append(data)

            except:
                continue

        published_candidates = []
        modified_candidates = []

        def scan_jsonld(obj):

            if isinstance(obj, dict):

                # Standard Schema.org fields
                for key in [
                    "datePublished",
                    "dateCreated",
                    "createdAt",
                    "created_at"
                ]:
                    if key in obj and obj[key]:
                        published_candidates.append(
                            (obj[key], f"JSON-LD:{key}")
                        )

                for key in [
                    "dateModified",
                    "modifiedAt",
                    "updatedAt",
                    "updated_at"
                ]:
                    if key in obj and obj[key]:
                        modified_candidates.append(
                            (obj[key], f"JSON-LD:{key}")
                        )

                # Recursively scan all values
                for value in obj.values():
                    scan_jsonld(value)

            elif isinstance(obj, list):

                for item in obj:
                    scan_jsonld(item)

        for obj in json_ld_objects:
            scan_jsonld(obj)


        # ----------------------------------------------------
        # B. META TAGS
        # ----------------------------------------------------

        meta_created_names = [
            "article:published_time",
            "datePublished",
            "dateCreated",
            "created",
            "created_at",
            "creation_date",
            "publish_date",
            "pubdate"
        ]

        meta_updated_names = [
            "article:modified_time",
            "dateModified",
            "dateModified",
            "modified",
            "updated",
            "updated_at",
            "last-modified",
            "last_modified",
            "modified_date",
            "update_date"
        ]

        for meta in soup.find_all("meta"):

            name = (
                meta.get("property")
                or meta.get("name")
                or meta.get("itemprop")
                or ""
            ).lower().strip()

            content = meta.get("content")

            if not content:
                continue

            if name in [x.lower() for x in meta_created_names]:
                published_candidates.append(
                    (content, f"META:{name}")
                )

            if name in [x.lower() for x in meta_updated_names]:
                modified_candidates.append(
                    (content, f"META:{name}")
                )


        # ----------------------------------------------------
        # C. VOLPAY / CMS SOURCE CODE
        # ----------------------------------------------------
        #
        # This is particularly important for the Volopay
        # source code you provided.
        #
        # Example:
        #
        # created_at:"2024-04-05T13:24:00.670Z"
        # updated_at:"2024-12-06T07:34:13.512Z"
        #
        # ----------------------------------------------------

        created_patterns = [
            r'created_at\s*[:=]\s*["\']([^"\']+)["\']',
            r'"created_at"\s*:\s*"([^"]+)"',
            r"'created_at'\s*:\s*'([^']+)'",
            r'createdAt\s*[:=]\s*["\']([^"\']+)["\']',
            r'"createdAt"\s*:\s*"([^"]+)"',
        ]

        updated_patterns = [
            r'updated_at\s*[:=]\s*["\']([^"\']+)["\']',
            r'"updated_at"\s*:\s*"([^"]+)"',
            r"'updated_at'\s*:\s*'([^']+)'",
            r'updatedAt\s*[:=]\s*["\']([^"\']+)["\']',
            r'"updatedAt"\s*:\s*"([^"]+)"',
        ]


        for pattern in created_patterns:

            matches = re.findall(
                pattern,
                source,
                flags=re.IGNORECASE
            )

            for match in matches:
                published_candidates.append(
                    (match, "SOURCE:created_at")
                )


        for pattern in updated_patterns:

            matches = re.findall(
                pattern,
                source,
                flags=re.IGNORECASE
            )

            for match in matches:
                modified_candidates.append(
                    (match, "SOURCE:updated_at")
                )


        # ----------------------------------------------------
        # D. PICK BEST DATE
        # ----------------------------------------------------
        #
        # Priority:
        #
        # Created:
        #   1. SOURCE:created_at
        #   2. JSON-LD
        #   3. META
        #
        # Updated:
        #   1. SOURCE:updated_at
        #   2. JSON-LD
        #   3. META
        #
        # For updated date, if multiple dates are available,
        # choose the latest valid date.
        # ----------------------------------------------------

        def select_date(candidates, preferred_sources):

            parsed = []

            for value, source_name in candidates:

                dt = parse_date_value(value)

                if dt:
                    parsed.append(
                        (dt, source_name)
                    )

            if not parsed:
                return None, ""

            # First try preferred source types
            for preferred in preferred_sources:

                preferred_matches = [
                    item
                    for item in parsed
                    if preferred.lower() in item[1].lower()
                ]

                if preferred_matches:

                    # Latest date within preferred source
                    return max(
                        preferred_matches,
                        key=lambda x: x[0]
                    )

            # Otherwise use latest available date
            return max(
                parsed,
                key=lambda x: x[0]
            )


        created_dt, created_source = select_date(
            published_candidates,
            [
                "SOURCE:created_at",
                "JSON-LD",
                "META"
            ]
        )

        updated_dt, updated_source = select_date(
            modified_candidates,
            [
                "SOURCE:updated_at",
                "JSON-LD",
                "META"
            ]
        )


        result["created_datetime"] = created_dt
        result["updated_datetime"] = updated_dt
        result["created_source"] = created_source
        result["updated_source"] = updated_source

        return result

    except Exception as e:

        result["error"] = str(e)

        return result


# ------------------------------------------------------------
# 6. GSC FUNCTION
# ------------------------------------------------------------

def get_gsc_queries_for_url(
    webmasters_service,
    site_url,
    page_url,
    start_date,
    end_date,
    max_queries=10
):

    request = {
        "startDate": start_date,
        "endDate": end_date,

        "dimensions": [
            "query"
        ],

        "dimensionFilterGroups": [
            {
                "filters": [
                    {
                        "dimension": "page",
                        "operator": "equals",
                        "expression": page_url
                    }
                ]
            }
        ],

        # Get enough rows to confidently select the top 10.
        "rowLimit": 25000,

        "startRow": 0
    }


    response = (
        webmasters_service
        .searchanalytics()
        .query(
            siteUrl=site_url,
            body=request
        )
        .execute()
    )


    rows = response.get("rows", [])

    results = []

    for row in rows:

        query = row.get("keys", [""])[0]

        impressions = row.get(
            "impressions",
            0
        )

        clicks = row.get(
            "clicks",
            0
        )

        ctr = row.get(
            "ctr",
            0
        )

        position = row.get(
            "position",
            0
        )

        results.append({
            "query": query,
            "impressions": impressions,
            "clicks": clicks,
            "ctr": ctr,
            "position": position
        })


    # --------------------------------------------------------
    # Sort by impressions HIGH -> LOW
    # --------------------------------------------------------

    results.sort(
        key=lambda x: x["impressions"],
        reverse=True
    )


    # --------------------------------------------------------
    # Keep maximum 10
    # --------------------------------------------------------

    return results[:max_queries]


# ------------------------------------------------------------
# 7. PROCESS ALL URLS
# ------------------------------------------------------------

all_rows = []

print("=" * 80)
print("STARTING REPORT")
print("=" * 80)

for index, url in enumerate(urls, start=1):

    print()
    print(f"[{index}/{len(urls)}] Processing:")
    print(url)


    # --------------------------------------------------------
    # Scrape created / updated dates
    # --------------------------------------------------------

    print("  → Reading page source...")

    page_dates = extract_page_dates(url)


    if page_dates["error"]:
        print(
            f"  ⚠ Date scraping issue: "
            f"{page_dates['error']}"
        )


    created_date = format_date_only(
        page_dates["created_datetime"]
    )

    updated_date = format_date_only(
        page_dates["updated_datetime"]
    )


    print(
        f"  → Created: "
        f"{created_date or 'Not found'}"
    )

    print(
        f"  → Updated: "
        f"{updated_date or 'Not found'}"
    )


    # --------------------------------------------------------
    # GSC
    # --------------------------------------------------------

    print("  → Fetching GSC queries...")

    try:

        gsc_queries = get_gsc_queries_for_url(
            webmasters_service=webmasters_service,
            site_url=SITE_URL,
            page_url=url,
            start_date=START_DATE,
            end_date=END_DATE,
            max_queries=MAX_QUERIES_PER_URL
        )

    except Exception as e:

        print(
            f"  ⚠ GSC error: {e}"
        )

        gsc_queries = []


    print(
        f"  → Queries found: "
        f"{len(gsc_queries)}"
    )


    # --------------------------------------------------------
    # Add rows
    # --------------------------------------------------------

    for item in gsc_queries:

        all_rows.append({

            "URL": url,

            "Created Date": created_date,

            "Updated Date": updated_date,

            "Queries Per URL": item["query"],

            "Impression": item["impressions"]

        })


# ------------------------------------------------------------
# 8. CREATE DATAFRAME
# ------------------------------------------------------------

df = pd.DataFrame(
    all_rows,
    columns=[
        "URL",
        "Created Date",
        "Updated Date",
        "Queries Per URL",
        "Impression"
    ]
)


# ------------------------------------------------------------
# 9. SAVE CSV
# ------------------------------------------------------------

df.to_csv(
    CSV_OUTPUT,
    index=False,
    encoding="utf-8-sig"
)

print()
print(f"CSV created: {CSV_OUTPUT}")


# ------------------------------------------------------------
# 10. CREATE EXCEL FILE
# ------------------------------------------------------------

with pd.ExcelWriter(
    XLSX_OUTPUT,
    engine="openpyxl"
) as writer:

    df.to_excel(
        writer,
        sheet_name="GSC Report",
        index=False
    )


# ------------------------------------------------------------
# 11. FORMAT EXCEL
# ------------------------------------------------------------

wb = load_workbook(
    XLSX_OUTPUT
)

ws = wb["GSC Report"]


# ------------------------------------------------------------
# Header formatting
# ------------------------------------------------------------

header_fill = PatternFill(
    fill_type="solid",
    fgColor="1F4E78"
)

header_font = Font(
    color="FFFFFF",
    bold=True
)

header_alignment = Alignment(
    horizontal="center",
    vertical="center"
)


for cell in ws[1]:

    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = header_alignment


# ------------------------------------------------------------
# Freeze header
# ------------------------------------------------------------

ws.freeze_panes = "A2"


# ------------------------------------------------------------
# Column widths
# ------------------------------------------------------------

column_widths = {
    "A": 65,
    "B": 20,
    "C": 20,
    "D": 55,
    "E": 15
}

for column, width in column_widths.items():

    ws.column_dimensions[column].width = width


# ------------------------------------------------------------
# Alignment
# ------------------------------------------------------------

for row in ws.iter_rows(
    min_row=2
):

    for cell in row:

        cell.alignment = Alignment(
            vertical="center",
            wrap_text=True
        )


# ------------------------------------------------------------
# 12. MERGE URL CELLS
# ------------------------------------------------------------
#
# Example:
#
# URL A
# URL A
# URL A
# URL A
#
# becomes:
#
# [        URL A        ]
# [                     ]
# [                     ]
# [                     ]
#
# ------------------------------------------------------------

current_row = 2

while current_row <= ws.max_row:

    url_value = ws.cell(
        current_row,
        1
    ).value

    start_row = current_row

    while (
        current_row + 1 <= ws.max_row
        and
        ws.cell(
            current_row + 1,
            1
        ).value == url_value
    ):

        current_row += 1

    end_row = current_row

    # Merge only when there are multiple rows
    if end_row > start_row:

        ws.merge_cells(
            start_row=start_row,
            start_column=1,
            end_row=end_row,
            end_column=1
        )

    # Vertically center merged URL
    ws.cell(
        start_row,
        1
    ).alignment = Alignment(
        vertical="center",
        horizontal="left",
        wrap_text=True
    )

    current_row += 1


# ------------------------------------------------------------
# 13. FIND DUPLICATE QUERIES ACROSS MULTIPLE URLS
# ------------------------------------------------------------

# A query should be red ONLY when it appears under
# more than one different URL.

query_to_urls = {}

for row in all_rows:

    query = row["Queries Per URL"]
    url = row["URL"]

    if not query:
        continue

    if query not in query_to_urls:
        query_to_urls[query] = set()

    query_to_urls[query].add(url)


duplicate_queries = {
    query
    for query, url_set in query_to_urls.items()
    if len(url_set) > 1
}


# ------------------------------------------------------------
# 14. HIGHLIGHT DUPLICATE QUERY CELLS RED
# ------------------------------------------------------------

red_fill = PatternFill(
    fill_type="solid",
    fgColor="FFC7CE"
)

red_font = Font(
    color="9C0006",
    bold=True
)


for row in range(2, ws.max_row + 1):

    query_cell = ws.cell(
        row,
        4
    )

    query = query_cell.value

    if query in duplicate_queries:

        query_cell.fill = red_fill
        query_cell.font = red_font


# ------------------------------------------------------------
# 15. ADD BORDERS
# ------------------------------------------------------------

thin_gray = Side(
    style="thin",
    color="D9E1F2"
)

border = Border(
    left=thin_gray,
    right=thin_gray,
    top=thin_gray,
    bottom=thin_gray
)

for row in ws.iter_rows():

    for cell in row:

        cell.border = border


# ------------------------------------------------------------
# 16. AUTO-FILTER
# ------------------------------------------------------------

ws.auto_filter.ref = ws.dimensions


# ------------------------------------------------------------
# 17. SAVE
# ------------------------------------------------------------

wb.save(
    XLSX_OUTPUT
)


# ------------------------------------------------------------
# 18. FINAL SUMMARY
# ------------------------------------------------------------

print()
print("=" * 80)
print("REPORT COMPLETE")
print("=" * 80)

print(
    f"URLs processed: {len(urls)}"
)

print(
    f"Total query rows: {len(df)}"
)

print(
    f"Duplicate queries across URLs: "
    f"{len(duplicate_queries)}"
)

print()
print(
    f"Excel file: {XLSX_OUTPUT}"
)

print(
    f"CSV file:   {CSV_OUTPUT}"
)

print("=" * 80)


# ------------------------------------------------------------
# 19. DOWNLOAD FILES IN GOOGLE COLAB
# ------------------------------------------------------------

from google.colab import files

print()
print("Downloading Excel report...")
files.download(XLSX_OUTPUT)

print()
print("Downloading CSV report...")
files.download(CSV_OUTPUT)
